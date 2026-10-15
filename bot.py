# ---------------------------------------------------------------------------
# Taxiatosh tumani 15-umumiy o'rta ta'lim maktabi
# O'quvchilar Kengashi saylovi: Telegram bot + Mini App serveri (bitta fayl)
#
#   1. /start bosilganda "Ariza topshirish" tugmasi chiqadi (Mini App ochiladi)
#   2. index.html ni internetga chiqaradi (Mini App shu)
#   3. Arizani tekshiradi, Supabase bazasiga saqlaydi va adminga
#      chiroyli Word (.docx) fayl yuboradi (ichida imzo ham bor)
# ---------------------------------------------------------------------------
import asyncio
import base64
from collections import Counter
import hashlib
import hmac
import json
import logging
import os
import re
import time
from datetime import datetime, timedelta, timezone
from html import escape as _html_escape
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qsl
 
import httpx
from aiogram import Bot, Dispatcher, F, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.filters import Command, CommandStart
from aiogram.types import (
    BufferedInputFile,
    CallbackQuery,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    MenuButtonWebApp,
    Message,
    WebAppInfo,
)
from aiohttp import web
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
 
logging.basicConfig(level=logging.INFO)
log = logging.getLogger("saylov")
 
 
def escape(text) -> str:
    """Telegram HTML uchun: faqat < > & almashtiriladi (apostrof o'z holicha qoladi)."""
    return _html_escape(str(text), quote=False)
 
# ----------------------------------------------------------------- SOZLAMALAR
BOT_TOKEN = os.environ["BOT_TOKEN"].strip()
SUPABASE_URL = os.environ["SUPABASE_URL"].strip().rstrip("/")
SUPABASE_KEY = os.environ["SUPABASE_KEY"].strip()
# Arizalar shu Telegram ID ga yuboriladi. Railway -> Variables -> ADMIN_ID
ADMIN_ID = int(os.getenv("ADMIN_ID", "1101690807").strip())
PORT = int(os.getenv("PORT", "8080"))
 
_domain = os.getenv("RAILWAY_PUBLIC_DOMAIN", "").strip()
WEBAPP_URL = os.getenv("WEBAPP_URL", "").strip().rstrip("/") or (
    f"https://{_domain}" if _domain else ""
)
 
SCHOOL = "Taxiatosh tumani 15-umumiy o'rta ta'lim maktabi"
HELP_LINE = "Agar botda muammo chiqsa, Madatga xabar bering."
TZ = timezone(timedelta(hours=5))  # Toshkent vaqti
 
GENERAL_COUNT = 10  # umumiy savollar soni (index.html dagi GENERAL bilan bir xil)
SPECIAL_COUNT = 4   # yo'nalish savollari soni (Kelajak liderligidan tashqari)
 
DIRECTIONS = {
    "leader": "Kelajak liderligi",
    "education": "Ta'lim koordinatori",
    "culture": "Madaniyat koordinatori",
    "it": "IT va sun'iy intellekt koordinatori",
    "sport": "Sport koordinatori",
    "eco": "Ekologiya koordinatori",
    "media": "Media koordinatori",
    "social": "Ijtimoiy tashabbuslar koordinatori",
}
 
# ------------------------------------------------------------- TEKSHIRUVLAR
WORD_RE = re.compile(r"[^\W\d_](?:[^\W\d_]|['‘’ʻʼ`´-])+")
PHONE_RE = re.compile(r"\+998\d{9}")
 
 
def check_name(raw: str):
    """To'liq Familiya + Ism bo'lsa, tozalangan nomni qaytaradi, aks holda None."""
    s = " ".join(str(raw).split())
    if not s or re.search(r"[.,\d_]", s):
        return None
    parts = s.split(" ")
    if not 2 <= len(parts) <= 4:
        return None
    for p in parts:
        if not WORD_RE.fullmatch(p):
            return None
        if sum(ch.isalpha() for ch in p) < 2:
            return None
    return " ".join(
        "-".join(seg[:1].upper() + seg[1:].lower() for seg in w.split("-"))
        for w in parts
    )
 
 
def name_key(name: str) -> str:
    """Ovozlarni sanash uchun kalit: bir odamning turlicha yozilishini birlashtiradi
    (katta-kichik harf, apostrof turi va so'zlar tartibi farqi hisobga olinmaydi)."""
    s = re.sub(r"[‘’ʻʼ`´]", "'", str(name).lower())
    return " ".join(sorted(s.split()))
 
 
def verify_init_data(init_data: str):
    """Telegram initData haqiqiyligini tekshiradi. To'g'ri bo'lsa user dict qaytaradi."""
    try:
        pairs = dict(parse_qsl(init_data, keep_blank_values=True))
        received = pairs.pop("hash", None)
        if not received:
            return None
        check = "\n".join(f"{k}={v}" for k, v in sorted(pairs.items()))
        secret = hmac.new(b"WebAppData", BOT_TOKEN.encode(), hashlib.sha256).digest()
        calc = hmac.new(secret, check.encode(), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(calc, received):
            return None
        if time.time() - int(pairs.get("auth_date", "0")) > 86400:
            return None
        user = json.loads(pairs["user"])
        return user if isinstance(user, dict) and "id" in user else None
    except Exception:
        return None
 
 
# ------------------------------------------------------------------ SUPABASE
def sb_headers(extra=None):
    h = {"apikey": SUPABASE_KEY, "Content-Type": "application/json"}
    if SUPABASE_KEY.startswith("eyJ"):  # eski service_role (JWT) kalit
        h["Authorization"] = f"Bearer {SUPABASE_KEY}"
    if extra:
        h.update(extra)
    return h
 
 
async def sb(method: str, path: str, **kw):
    headers = sb_headers(kw.pop("headers", None))
    async with httpx.AsyncClient(timeout=25) as client:
        return await client.request(
            method, f"{SUPABASE_URL}/rest/v1/{path}", headers=headers, **kw
        )
 
 
# --------------------------------------------------------------------- BOT
router = Router()
 
 
# ---------------------------------------------------- OVOZ BERISH SOZLAMASI
async def voting_open() -> bool:
    """Ovoz berish yoqilganmi? Jadval yo'q yoki xato bo'lsa, yoqilgan deb hisoblanadi."""
    try:
        r = await sb("GET", "settings?key=eq.voting_open&select=value")
        if r.status_code == 200:
            rows = r.json()
            if rows:
                return str(rows[0]["value"]).lower() == "true"
            return True
        log.warning("settings o'qib bo'lmadi: %s %s", r.status_code, r.text[:200])
    except Exception:
        log.exception("settings o'qishda xato")
    return True
 
 
async def set_voting(open_: bool) -> bool:
    try:
        r = await sb(
            "POST",
            "settings?on_conflict=key",
            json={"key": "voting_open", "value": "true" if open_ else "false"},
            headers={"Prefer": "resolution=merge-duplicates,return=minimal"},
        )
        if r.status_code in (200, 201, 204):
            return True
        log.error("settings yozib bo'lmadi: %s %s", r.status_code, r.text[:300])
    except Exception:
        log.exception("settings yozishda xato")
    return False
 
 
async def fetch_rows():
    r = await sb("GET", "applications?select=grade,full_name,vote_for&order=created_at.asc&limit=5000")
    if r.status_code != 200:
        raise RuntimeError(f"{r.status_code}: {r.text[:200]}")
    return r.json()
 
 
def format_results(rows: list, is_open: bool) -> str:
    """Kim necha ovoz olgani: umumiy va sinflar bo'yicha, foizda."""
    voters = [x for x in rows if x.get("vote_for")]
    total = len(voters)
    head = [
        "📊 <b>Ovoz berish natijalari</b>",
        "Holat: " + ("🟢 yoqilgan" if is_open else "🔴 to'xtatilgan"),
        f"Jami arizalar: {len(rows)}",
        f"Ovoz berganlar: {total}",
    ]
    if not total:
        return "\n".join(head + ["", "Hozircha ovoz yo'q."])
 
    def table(items, limit):
        groups = {}
        for x in items:
            groups.setdefault(name_key(x["vote_for"]), []).append(x["vote_for"])
        ranked = sorted(groups.values(), key=lambda v: (-len(v), Counter(v).most_common(1)[0][0]))
        n = len(items)
        out = []
        for i, names in enumerate(ranked[:limit], 1):
            shown = Counter(names).most_common(1)[0][0]
            out.append(f"{i}. {escape(shown)} — {len(names)} ta ({len(names) / n * 100:.1f}%)")
        if len(ranked) > limit:
            out.append(f"… yana {len(ranked) - limit} ta ism")
        return out
 
    lines = head + ["", "<b>Umumiy (hamma sinflar):</b>"] + table(voters, 30)
    for g in range(5, 12):
        items = [x for x in voters if x["grade"] == g]
        if items:
            lines += ["", f"<b>{g}-sinf o'quvchilari ({len(items)} ta ovoz):</b>"] + table(items, 5)
    return "\n".join(lines)
 
 
async def send_long(bot: Bot, chat_id: int, text: str, reply_markup=None):
    parts, chunk = [], ""
    for line in text.split("\n"):
        if len(chunk) + len(line) + 1 > 3800:
            parts.append(chunk)
            chunk = ""
        chunk += line + "\n"
    if chunk.strip():
        parts.append(chunk)
    for i, part in enumerate(parts):
        await bot.send_message(chat_id, part, reply_markup=reply_markup if i == len(parts) - 1 else None)
 
 
def admin_kb(is_open: bool) -> InlineKeyboardMarkup:
    if is_open:
        toggle = InlineKeyboardButton(text="🛑 Ovozni to'xtatish", callback_data="adm:stop")
    else:
        toggle = InlineKeyboardButton(text="▶️ Ovozni qayta yoqish", callback_data="adm:start")
    return InlineKeyboardMarkup(
        inline_keyboard=[[InlineKeyboardButton(text="📊 Natijalar", callback_data="adm:results")], [toggle]]
    )
 
 
async def send_results(bot: Bot, chat_id: int):
    is_open = await voting_open()
    try:
        rows = await fetch_rows()
    except Exception as e:
        await bot.send_message(
            chat_id,
            "⚠️ Natijalarni o'qib bo'lmadi. Supabase'da yangi supabase.sql ni ishga tushirganmisiz?\n"
            f"<code>{escape(str(e))}</code>",
        )
        return
    await send_long(bot, chat_id, format_results(rows, is_open), admin_kb(is_open))
 
 
@router.message(Command("admin"))
async def cmd_admin(m: Message, bot: Bot):
    if m.from_user.id != ADMIN_ID:
        return
    await send_results(bot, m.chat.id)
 
 
@router.callback_query(F.data.startswith("adm:"))
async def on_admin_button(cb: CallbackQuery, bot: Bot):
    if cb.from_user.id != ADMIN_ID:
        await cb.answer("Ruxsat yo'q", show_alert=True)
        return
    action = cb.data.split(":", 1)[1]
    chat_id = cb.from_user.id
    await cb.answer()
    if action == "results":
        await send_results(bot, chat_id)
    elif action == "stop":
        kb = InlineKeyboardMarkup(
            inline_keyboard=[[
                InlineKeyboardButton(text="✅ Ha, to'xtatish", callback_data="adm:stop_yes"),
                InlineKeyboardButton(text="❌ Yo'q", callback_data="adm:cancel"),
            ]]
        )
        await bot.send_message(
            chat_id,
            "🛑 <b>Ovoz berishni to'xtatmoqchimisiz?</b>\n\n"
            "Shundan keyin yangi arizalarda “Kimga ovoz berasiz?” savoli so'ralmaydi. "
            "Keyin qayta yoqish mumkin.",
            reply_markup=kb,
        )
    elif action == "stop_yes":
        if await set_voting(False):
            await bot.send_message(chat_id, "🛑 Ovoz berish <b>to'xtatildi</b>. Yakuniy natijalar:")
            await send_results(bot, chat_id)
        else:
            await bot.send_message(chat_id, "⚠️ Saqlab bo'lmadi. Supabase'da yangi supabase.sql ni ishga tushiring.")
    elif action == "start":
        if await set_voting(True):
            await bot.send_message(chat_id, "▶️ Ovoz berish <b>qayta yoqildi</b>.", reply_markup=admin_kb(True))
        else:
            await bot.send_message(chat_id, "⚠️ Saqlab bo'lmadi. Supabase'da yangi supabase.sql ni ishga tushiring.")
    elif action == "cancel":
        await bot.send_message(chat_id, "Bekor qilindi.", reply_markup=admin_kb(await voting_open()))
 
 
@router.message(CommandStart())
async def cmd_start(m: Message):
    if not WEBAPP_URL:
        await m.answer("⚠️ Bot hali to'liq sozlanmagan (WEBAPP_URL yo'q). Admin bilan bog'laning.")
        return
    kb = InlineKeyboardMarkup(
        inline_keyboard=[
            [InlineKeyboardButton(text="📝 Ariza topshirish", web_app=WebAppInfo(url=WEBAPP_URL))]
        ]
    )
    await m.answer(
        "👋 <b>Assalomu alaykum!</b>\n\n"
        f"🏫 {SCHOOL} <b>O'quvchilar Kengashi</b> "
        "saylovida ishtirok etish uchun ariza topshiring.\n\n"
        "Quyidagi tugmani bosing 👇 va barcha savollarga javob bering. "
        "Oxirida imzo qo'yasiz.\n\n"
        f"ℹ️ {HELP_LINE}",
        reply_markup=kb,
    )
    if m.from_user.id == ADMIN_ID:
        await m.answer(
            "🔐 <b>Admin panel</b>\nNatijalarni ko'rish va ovozni to'xtatish shu yerda.",
            reply_markup=admin_kb(await voting_open()),
        )
 
 
@router.message(Command("stats"))
async def cmd_stats(m: Message):
    if m.from_user.id != ADMIN_ID:
        return
    r = await sb(
        "GET",
        "applications?select=grade,full_name,direction&order=grade.asc,created_at.asc",
    )
    if r.status_code != 200:
        await m.answer(f"Bazadan o'qib bo'lmadi: {r.status_code}")
        return
    rows = r.json()
    if not rows:
        await m.answer("Hozircha ariza yo'q.")
        return
    lines = [f"📊 <b>Jami arizalar: {len(rows)}</b>"]
    for g in range(5, 12):
        items = [x for x in rows if x["grade"] == g]
        lines.append(f"\n<b>{g}-sinf</b>: {len(items)} ta")
        for x in items:
            lines.append(f"  • {escape(x['full_name'])} — {escape(x['direction'])}")
    chunk = ""
    for line in lines:
        if len(chunk) + len(line) > 3800:
            await m.answer(chunk)
            chunk = ""
        chunk += line + "\n"
    if chunk:
        await m.answer(chunk)
 
 
# ------------------------------------------------------------ WORD FAYL
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x55, 0x55)
 
 
def _shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)
 
 
def _shade_par(par, fill):
    p_pr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
 
 
def _par(doc, text="", bold=False, size=11, align=None, color=None, after=4, before=0, italic=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        if color is not None:
            r.font.color.rgb = color
    return p
 
 
def _heading(doc, text):
    p = _par(doc, text, bold=True, size=12, color=RGBColor(0xFF, 0xFF, 0xFF), after=6, before=12)
    _shade_par(p, "1F3A5F")
    p.paragraph_format.left_indent = Cm(0.1)
    return p
 
 
def build_docx(row: dict, sig_bytes: bytes, submitted: datetime) -> bytes:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
 
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
 
    C = WD_ALIGN_PARAGRAPH.CENTER
    _par(doc, SCHOOL.upper(), bold=True, size=12, align=C, color=GREY, after=2)
    _par(doc, "O'QUVCHILAR KENGASHI SAYLOVI", bold=True, size=16, align=C, color=NAVY, after=2)
    _par(doc, "A R I Z A", bold=True, size=20, align=C, color=NAVY, after=2)
    _par(
        doc,
        f"Ariza № {row.get('id', '-')}   •   {submitted:%d.%m.%Y, %H:%M}",
        size=10,
        align=C,
        color=GREY,
        after=10,
    )
 
    info = [
        ("Familiya, ism", row["full_name"]),
        ("Sinf", f"{row['grade']}-sinf"),
        ("Nomzodlik yo'nalishi", row["direction"]),
        ("Saylov bosqichi", "Maktab bosqichi"),
        ("Aymoq", row["region"]),
        ("Ta'lim muassasasi", row["institution"]),
        ("Tug'ilgan sana", datetime.strptime(row["birth_date"], "%Y-%m-%d").strftime("%d.%m.%Y")),
        ("Telefon raqam", row["phone"]),
        ("Sardor bo'lolmasa, ovoz beradi", row.get("vote_for") or "—"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in info:
        cells = table.add_row().cells
        cells[0].width, cells[1].width = Cm(5.2), Cm(11.8)
        _shade_cell(cells[0], "E8EEF6")
        lp = cells[0].paragraphs[0]
        lr = lp.add_run(label)
        lr.bold = True
        lr.font.size = Pt(10.5)
        vp = cells[1].paragraphs[0]
        vr = vp.add_run(str(value))
        vr.font.size = Pt(11)
        if label in ("Familiya, ism", "Nomzodlik yo'nalishi"):
            vr.bold = True
 
    section = None
    counter = 0
    for a in row["answers"]:
        if a["section"] != section:
            section = a["section"]
            counter = 0
            title = "UMUMIY SAVOLLAR" if section == "Umumiy savol" else f"YO'NALISH SAVOLLARI — {row['direction'].upper()}"
            _heading(doc, title)
        counter += 1
        _par(doc, f"{counter}. {a['q']}", bold=True, size=11, color=NAVY, after=2, before=4)
        ans = _par(doc, a["a"], size=11, after=6)
        ans.paragraph_format.left_indent = Cm(0.6)
 
    _heading(doc, "NOMZOD MAJBURIYATLARI")
    _par(
        doc,
        "Men ushbu arizani topshirish orqali O'quvchilar Kengashi faoliyatining ixtiyoriyligi, tenglik, "
        "ochiqlik va oshkoralik, hurmat, mas'uliyat, tayyorlik va natijadorlik tamoyillariga amal qilishga "
        "roziman. Saylov vaqtida boshqa ishtirokchilarning huquq va manfaatlarini hurmat qilish, ovoz berish "
        "vaqtida bosim o'tkazmaslik hamda saylov natijalarini hurmat qilish majburiyatini olaman.",
        size=10.5,
        after=8,
    )
 
    _par(doc, f"Nomzod: {row['full_name']}", bold=True, after=2)
    _par(doc, "Qo'li:", bold=True, after=2)
    sig_par = _par(doc, "", after=2)
    sig_par.add_run().add_picture(BytesIO(sig_bytes), width=Cm(5.5))
    _par(doc, f"Sana: {submitted:%d.%m.%Y}", bold=True, after=0)
 
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
 
 
def build_text(row: dict, submitted: datetime) -> str:
    """Word fayl yaratib bo'lmasa, zaxira sifatida oddiy matn."""
    L = [
        f"{SCHOOL.upper()}",
        "O'QUVCHILAR KENGASHI SAYLOVI — ARIZA",
        "=" * 40,
        f"Ariza raqami: {row.get('id', '-')}",
        f"Vaqt: {submitted:%d.%m.%Y %H:%M}",
        f"Familiya, ism: {row['full_name']}",
        f"Sinf: {row['grade']}-sinf",
        f"Yo'nalish: {row['direction']}",
        f"Aymoq: {row['region']}",
        f"Muassasa: {row['institution']}",
        f"Tug'ilgan sana: {row['birth_date']}",
        f"Telefon: {row['phone']}",
        f"Sardor bo'lolmasa, ovoz beradi: {row.get('vote_for') or '-'}",
        "",
    ]
    for i, a in enumerate(row["answers"], 1):
        L += [f"{i}. [{a['section']}] {a['q']}", a["a"], ""]
    return "\n".join(L)
 
 
async def notify_admin(bot: Bot, row: dict, sig_bytes: bytes) -> bool:
    submitted = datetime.now(TZ)
    caption = (
        f"🆕 <b>Yangi ariza #{row.get('id', '')}</b>\n\n"
        f"👤 <b>{escape(row['full_name'])}</b>\n"
        f"🏫 {row['grade']}-sinf\n"
        f"🎯 {escape(row['direction'])}\n"
        f"📞 {escape(row['phone'])}\n"
        f"🗳 Ovoz beradi: {escape(row['vote_for']) if row.get('vote_for') else '—'}\n\n"
        f"📄 To'liq ariza (imzo bilan) faylda 👇"
    )
    safe = re.sub(r"\W+", "_", row["full_name"]).strip("_") or "ariza"
    base = f"ariza_{row['grade']}sinf_{safe}"
    try:
        try:
            data = build_docx(row, sig_bytes, submitted)
            fname = base + ".docx"
        except Exception:
            log.exception("Word fayl yaratib bo'lmadi, matn yuboriladi")
            data = build_text(row, submitted).encode("utf-8")
            fname = base + ".txt"
        await bot.send_document(
            ADMIN_ID, BufferedInputFile(data, fname), caption=caption, reply_markup=admin_kb(await voting_open())
        )
        return True
    except Exception:
        log.exception("Adminga yuborib bo'lmadi. Admin botga /start bosganmi? ADMIN_ID to'g'rimi?")
        return False
 
 
def bad(msg: str, status: int = 400):
    return web.json_response({"ok": False, "error": msg}, status=status)
 
 
async def api_submit(request: web.Request):
    bot: Bot = request.app["bot"]
    try:
        body = await request.json()
    except Exception:
        return bad("So'rov noto'g'ri.")
 
    user = verify_init_data(str(body.get("initData", "")))
    if not user:
        return bad("Ilovani Telegram bot orqali oching (/start bosing).", 403)
 
    d = body.get("data")
    if not isinstance(d, dict):
        return bad("Ma'lumotlar noto'g'ri.")
 
    grade = d.get("cls")
    if not isinstance(grade, int) or isinstance(grade, bool) or not 5 <= grade <= 11:
        return bad("Sinf noto'g'ri tanlangan.")
 
    name = check_name(d.get("fullName", ""))
    if not name:
        return bad("Familiya va ismni to'liq yozing, qisqartirish yoki faqat ism yozish taqiqlanadi.")
 
    vote_for = None
    if await voting_open():
        vote_for = check_name(d.get("voteFor", ""))
        if not vote_for:
            return bad("Kimga ovoz berishingizni to'liq familiya va ism bilan yozing (masalan: Qazaqbayev Jamshid).")
        if name_key(vote_for) == name_key(name):
            return bad("O'zingizga ovoz bera olmaysiz. Boshqa nomzodni yozing.")
 
    dir_id = d.get("direction")
    if dir_id not in DIRECTIONS:
        return bad("Yo'nalish noto'g'ri tanlangan.")
 
    region = str(d.get("region", "")).strip()[:200]
    institution = str(d.get("institution", "")).strip()[:200]
    if len(region) < 3 or len(institution) < 3:
        return bad("Aymoq va muassasa nomini yozing.")
 
    try:
        birth = datetime.strptime(str(d.get("birth", "")), "%Y-%m-%d")
        if not 2005 <= birth.year <= 2017:
            raise ValueError
    except ValueError:
        return bad("Tug'ilgan sana noto'g'ri.")
 
    phone = str(d.get("phone", "")).strip()
    if not PHONE_RE.fullmatch(phone):
        return bad("Telefon raqami noto'g'ri.")
 
    if d.get("agree") is not True:
        return bad("Majburiyatlarga rozilik kerak.")
 
    answers = []
    counts = {}
    for section, key in (("Umumiy savol", "general"), ("Yo'nalish savoli", "special")):
        items = d.get(key) or []
        if not isinstance(items, list) or len(items) > 20:
            return bad("Javoblar noto'g'ri.")
        counts[key] = len(items)
        for it in items:
            if not isinstance(it, dict):
                return bad("Javoblar noto'g'ri.")
            q = str(it.get("q", "")).strip()[:300]
            a = str(it.get("a", "")).strip()
            if len(a) < 10 or len(a) > 1500:
                return bad("Har bir javob kamida 10 ta belgidan iborat bo'lishi kerak.")
            answers.append({"section": section, "q": q, "a": a})
    need_special = 0 if dir_id == "leader" else SPECIAL_COUNT
    if counts["general"] != GENERAL_COUNT or counts["special"] != need_special:
        return bad("Savollar soni to'liq emas. Sahifani yangilab, qaytadan urinib ko'ring.")
 
    sig = str(d.get("signature", ""))
    if not sig.startswith("data:image/png;base64,") or len(sig) > 600_000:
        return bad("Imzo noto'g'ri.")
    try:
        sig_bytes = base64.b64decode(sig.split(",", 1)[1])
    except Exception:
        return bad("Imzo noto'g'ri.")
 
    row = {
        "telegram_id": int(user["id"]),
        "telegram_username": user.get("username"),
        "grade": grade,
        "full_name": name,
        "direction": DIRECTIONS[dir_id],
        "region": region,
        "institution": institution,
        "birth_date": birth.strftime("%Y-%m-%d"),
        "phone": phone,
        "answers": answers,
        "vote_for": vote_for,
        "signature": sig,
        "agreed": True,
    }
 
    try:
        r = await sb("POST", "applications", json=row, headers={"Prefer": "return=representation"})
    except Exception:
        log.exception("Supabase bilan aloqa yo'q")
        return bad("Server bazaga ulana olmadi. Birozdan keyin qayta urinib ko'ring.", 502)
 
    if r.status_code == 409:
        return bad("Siz allaqachon ariza topshirgansiz. Har bir o'quvchi faqat bitta ariza topshiradi.")
    if r.status_code not in (200, 201):
        log.error("Supabase xato %s: %s", r.status_code, r.text[:500])
        return bad("Arizani saqlashda xatolik. Keyinroq urinib ko'ring.", 500)
 
    try:
        row["id"] = r.json()[0]["id"]
    except Exception:
        row["id"] = "-"
 
    admin_ok = await notify_admin(bot, row, sig_bytes)
 
    try:
        await bot.send_message(
            row["telegram_id"],
            "✅ <b>Arizangiz qabul qilindi!</b>\n\n"
            f"👤 {escape(name)}\n🏫 {grade}-sinf\n🎯 {escape(DIRECTIONS[dir_id])}\n\n"
            "Omad! 🍀\n\n"
            f"ℹ️ {HELP_LINE}",
        )
    except Exception:
        log.warning("O'quvchiga tasdiq xabari yuborilmadi")
 
    return web.json_response({"ok": True, "admin_notified": admin_ok})
 
 
async def api_status(request: web.Request):
    return web.json_response({"voting": await voting_open()}, headers={"Cache-Control": "no-store"})
 
 
async def index(request: web.Request):
    return web.FileResponse(
        Path(__file__).with_name("index.html"), headers={"Cache-Control": "no-store"}
    )
 
 
async def health(request: web.Request):
    return web.Response(text="ok")
 
 
# -------------------------------------------------------------------- MAIN
async def main():
    bot = Bot(BOT_TOKEN, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
    dp = Dispatcher()
    dp.include_router(router)
 
    app = web.Application(client_max_size=2 * 1024 * 1024)
    app["bot"] = bot
    app.router.add_get("/", index)
    app.router.add_get("/health", health)
    app.router.add_get("/api/status", api_status)
    app.router.add_post("/api/submit", api_submit)
 
    runner = web.AppRunner(app)
    await runner.setup()
    await web.TCPSite(runner, "0.0.0.0", PORT).start()
    log.info("Veb-server %s portda ishga tushdi. WEBAPP_URL=%s", PORT, WEBAPP_URL or "(yo'q)")
 
    await bot.delete_webhook(drop_pending_updates=True)
    if WEBAPP_URL:
        try:
            await bot.set_chat_menu_button(
                menu_button=MenuButtonWebApp(text="Ariza topshirish", web_app=WebAppInfo(url=WEBAPP_URL))
            )
        except Exception:
            log.exception("Menyu tugmasini o'rnatib bo'lmadi")
 
    await dp.start_polling(bot, allowed_updates=["message", "callback_query"])
 
 
if __name__ == "__main__":
    asyncio.run(main())